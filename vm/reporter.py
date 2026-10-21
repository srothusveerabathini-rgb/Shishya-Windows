#!/usr/bin/env python3
"""
Shishya — turn handed-over files into a report with graphs. Runs INSIDE the sandbox.

    python3 reporter.py /home/clone/shared/jobs/<stamp>

The principal's-report flow: drop the school's performance workbook on the
panel, press Make a report, and get back averages per class, the trend across
months, and the best and worst performers — as charts in one HTML file, plus
each chart as its own PNG.

Deliberately NOT a free-form job. The report's shape is written code — which
columns exist, group, average, plot — because the shape of a good report does
not vary and the model writing matplotlib from scratch fails in ways a person
has to debug live. The model gets exactly one bounded decision per table:
which column names the groups and which columns hold the scores. If it fails
or answers nonsense, heuristics decide instead and the report still arrives.
Same rule Findings learned: the model where the shape varies, code where it
does not.
"""

import base64
import datetime as dt
import io
import json
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent / "shell"))
import common      # noqa: E402

import matplotlib                      # noqa: E402
matplotlib.use("Agg")                  # no display needed, no display used
import matplotlib.pyplot as plt        # noqa: E402
import pandas as pd                    # noqa: E402

# One look for every chart, set once. Close to the panel's language: quiet
# ground, ink bars, no palette to fight the projector.
plt.rcParams.update({
    "figure.facecolor": "white", "axes.facecolor": "white",
    "axes.edgecolor": "#444", "axes.labelcolor": "#222",
    "xtick.color": "#444", "ytick.color": "#444",
    "axes.grid": True, "grid.color": "#e5e5e5", "grid.linewidth": .6,
    "font.size": 10, "axes.titlesize": 12, "axes.titleweight": "bold",
    "figure.dpi": 110,
})
INK = "#1f2430"


def put_status(job, **fields):
    """Same contract as worker.py: whole file, renamed into place."""
    fields.setdefault("t", round(time.time(), 1))
    tmp = job / ".status.tmp"
    tmp.write_text(json.dumps(fields, ensure_ascii=False), encoding="utf-8")
    tmp.replace(job / "status.json")


# ─── reading whatever arrived ───────────────────────────────────────────────

def header_row(ws, look=14):
    """Widest row near the top — same rule chat uses, kept dependency-free
    here so a report never fails because a chat import did."""
    best, widest = 0, 0
    for n, row in enumerate(ws.iter_rows(min_row=1, max_row=look, values_only=True)):
        filled = sum(1 for c in row if c is not None and str(c).strip())
        if filled > widest:
            best, widest = n, filled
    return best


def tables_from(path):
    """Every table a file contains, as (name, DataFrame). Word tables count —
    a principal's report is as often a docx as a workbook."""
    suffix = path.suffix.lower()
    out = []
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            at = header_row(ws)
            wb2 = pd.read_excel(path, sheet_name=ws.title, header=at)
            if wb2.dropna(how="all").shape[0] >= 2:
                out.append((f"{path.stem} — {ws.title}", wb2))
        wb.close()
    elif suffix in (".csv", ".tsv"):
        out.append((path.stem, pd.read_csv(path, sep=None, engine="python")))
    elif suffix == ".docx":
        from docx import Document
        for i, table in enumerate(Document(str(path)).tables, 1):
            rows = [[c.text.strip() for c in r.cells] for r in table.rows]
            if len(rows) >= 3:
                frame = pd.DataFrame(rows[1:], columns=rows[0])
                frame = frame.apply(pd.to_numeric, errors="ignore")
                out.append((f"{path.stem} — table {i}", frame))
    return out


# ─── deciding what the columns mean ─────────────────────────────────────────

def looks_dated(name):
    if isinstance(name, (dt.datetime, dt.date)):
        return True
    s = str(name).strip().lower()
    months = ("jan", "feb", "mar", "apr", "may", "jun", "jul",
              "aug", "sep", "oct", "nov", "dec")
    return s.startswith(months) or (len(s) in (6, 7) and s[:4].isdigit())


def by_heuristic(frame):
    """The fallback that always answers. Group = the leftmost text column
    whose values repeat or name people; scores = every numeric column that
    is not a serial number; months = columns whose headers read as dates."""
    months = [c for c in frame.columns if looks_dated(c)]
    numeric = [c for c in frame.columns
               if c not in months and pd.to_numeric(frame[c], errors="coerce").notna().mean() > .7]
    numeric = [c for c in numeric
               if not str(c).strip().lower() in ("sno", "s.no", "sl no", "roll no", "id", "sr no")
               and frame[c].nunique() > 1]
    # Prefer the column that names things over the one that numbers them:
    # a class sheet has Roll first and Name second, and "best Roll is 9B07"
    # is a finding nobody can read out loud at assembly.
    NAMED = ("name", "class", "section", "student", "product", "item",
             "supplier", "category", "subject", "school", "branch")
    candidates = []
    for c in frame.columns:
        if c in months or c in numeric:
            continue
        vals = frame[c].dropna().astype(str)
        if len(vals) and vals.nunique() >= 2:
            candidates.append(c)
    group = next((c for c in candidates
                  if any(w in str(c).strip().lower() for w in NAMED)),
                 candidates[0] if candidates else None)
    return {"group": group, "scores": numeric[:8], "months": months}


def by_model(name, frame):
    """One bounded question, then done. A wrong answer here costs nothing —
    anything that fails validation falls back to the heuristic."""
    cols = [str(c) for c in frame.columns][:40]
    schema = {"type": "object",
              "properties": {"group": {"type": "string"},
                             "scores": {"type": "array", "items": {"type": "string"}}},
              "required": ["group", "scores"]}
    try:
        raw = common.ask_job("answers",
                             f"Table {name!r} has columns: {', '.join(cols)}\n"
                             f"Which single column names the groups being compared "
                             f"(class, section, student, product...), and which "
                             f"columns hold the numeric scores or marks to average? "
                             f"Use exact column names.",
                             schema=schema, budget=300, timeout=90)
        got = json.loads(raw)
        namemap = {str(c): c for c in frame.columns}
        group = namemap.get(got.get("group"))
        scores = [namemap[s] for s in got.get("scores", []) if s in namemap]
        scores = [s for s in scores
                  if pd.to_numeric(frame[s], errors="coerce").notna().mean() > .5]
        if group is not None and scores:
            return {"group": group, "scores": scores[:8],
                    "months": [c for c in frame.columns if looks_dated(c)]}
    except Exception:
        pass
    return None


# ─── drawing ────────────────────────────────────────────────────────────────

def save(fig, out_dir, name, charts):
    fig.tight_layout()
    path = out_dir / name
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    buf = io.BytesIO()
    # Re-encode from disk so the HTML embeds exactly the file offered below it.
    buf.write(path.read_bytes())
    charts.append((name, base64.b64encode(buf.getvalue()).decode()))


def chart_table(name, frame, roles, out_dir, charts, findings):
    group, scores, months = roles["group"], roles["scores"], roles["months"]
    stem = "".join(ch if ch.isalnum() else "-" for ch in name)[:40].strip("-")

    for c in scores:
        frame[c] = pd.to_numeric(frame[c], errors="coerce")

    # 1 — the averages. One bar per group per score column.
    if group and scores:
        means = frame.groupby(group, dropna=True)[scores].mean(numeric_only=True)
        means = means.dropna(how="all")
        if 2 <= len(means) <= 30:
            width = min(12, max(6, .55 * len(means) * max(1, len(scores))))
            fig, ax = plt.subplots(figsize=(width, 3.6))
            means.plot.bar(ax=ax, color=[INK, "#5b6472", "#9aa1ac", "#c2c7cf"][:len(scores)],
                           edgecolor="white", linewidth=.5)
            ax.set_title(f"Average by {group}")
            ax.set_xlabel("")
            ax.legend(fontsize=8, frameon=False)
            save(fig, out_dir, f"{stem}-averages.png", charts)
            values = frame[scores].stack()
            if len(scores) == 1 or (values.min() >= 0 and values.max() <= 100):
                # Marks out of 100 average meaningfully across subjects.
                overall = means.mean(axis=1).sort_values()
                findings.append(f"{name}: best {group} is {overall.index[-1]} "
                                f"({overall.iloc[-1]:.1f}), weakest is {overall.index[0]} "
                                f"({overall.iloc[0]:.1f}); overall average {overall.mean():.1f}.")
            else:
                # Sales and Profit do not: their average is a number about
                # nothing ("best month 960" from sales 1500 and profit 420).
                for c in scores:
                    col = means[c].dropna().sort_values()
                    if len(col) >= 2:
                        findings.append(f"{name}: {c} is highest for {col.index[-1]} "
                                        f"({col.iloc[-1]:,.1f}) and lowest for {col.index[0]} "
                                        f"({col.iloc[0]:,.1f}).")

    # 2 — the trend. Month columns become one line, per group when few enough.
    if months:
        monthly = frame[months].apply(pd.to_numeric, errors="coerce")
        labels = [c.strftime("%b %y") if isinstance(c, (dt.datetime, dt.date))
                  else str(c)[:8] for c in months]
        fig, ax = plt.subplots(figsize=(7, 3.4))
        if group and frame[group].nunique() <= 8:
            for key, part in frame.groupby(group):
                ax.plot(labels, part[months].apply(pd.to_numeric, errors="coerce").mean(),
                        marker="o", markersize=3, linewidth=1.4, label=str(key)[:20])
            ax.legend(fontsize=8, frameon=False)
        else:
            ax.plot(labels, monthly.mean(), marker="o", color=INK, linewidth=1.6)
        ax.set_title("Across the months")
        ax.tick_params(axis="x", rotation=45)
        save(fig, out_dir, f"{stem}-trend.png", charts)
        seq = monthly.mean()
        if seq.notna().sum() >= 2:
            first, last = seq.dropna().iloc[0], seq.dropna().iloc[-1]
            word = "rose" if last > first else "fell" if last < first else "held steady at"
            findings.append(f"{name}: the average {word} from {first:.1f} to {last:.1f} "
                            f"across {seq.notna().sum()} months.")

    # 3 — the ends. Top and bottom five rows by their own average.
    if scores and len(frame) > 10:
        label = group or frame.columns[0]
        rowmean = frame[scores].mean(axis=1)
        picked = pd.concat([rowmean.nlargest(5), rowmean.nsmallest(5)]).sort_values()
        names = frame.loc[picked.index, label].astype(str).str.slice(0, 24)
        fig, ax = plt.subplots(figsize=(6.5, 3.6))
        ax.barh(names, picked.values,
                color=[INK if v >= rowmean.median() else "#9aa1ac" for v in picked.values])
        ax.set_title("Top 5 and bottom 5")
        save(fig, out_dir, f"{stem}-ends.png", charts)


# ─── the page ───────────────────────────────────────────────────────────────

def provenance(job):
    """One line, quoting the host's own record of the machine at the time.

    Not the sandbox's opinion of itself: it cannot check, and a deliverable
    that vouches for its own conditions is worth nothing. Says "not recorded"
    rather than guessing, because an unverified claim of a locked network is
    exactly the failure this is here to prevent.
    """
    try:
        spec = json.loads((job / "job.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return "Made on this machine. Network state at the time was not recorded."
    net = spec.get("egress") or {}
    locked = net.get("locked")
    where = ("with the sandbox's network locked — nothing could leave this "
             "computer while this was made" if locked is True else
             "with the sandbox's network OPEN — " + (net.get("detail") or "")
             if locked is False else
             "with the network state unrecorded")
    return (f"Made by Shishya on this machine, {where}. "
            f"Model: {spec.get('model', '?')}.")


def html_report(title, sections, findings, skipped, stamp=""):
    css = ("body{font-family:-apple-system,'Segoe UI',sans-serif;max-width:860px;"
           "margin:0 auto;padding:28px;color:#1a1c1e;background:#fcfcfa}"
           "h1{font-size:26px;border-bottom:2px solid #1a1c1e;padding-bottom:8px}"
           "h2{font-size:16px;margin-top:30px}img{max-width:100%;border:1px solid #ddd;"
           "margin:8px 0}li{margin:5px 0}.dim{color:#71767c;font-size:13px}")
    parts = [f"<!doctype html><meta charset='utf-8'><title>{title}</title>"
             f"<style>{css}</style><h1>{title}</h1>",
             f"<p class='dim'>{time.strftime('%d %b %Y, %H:%M')} — {stamp} "
             f"Every chart is also beside this file as its own PNG.</p>"]
    if findings:
        parts.append("<h2>What stands out</h2><ul>")
        parts += [f"<li>{f}</li>" for f in findings[:10]]
        parts.append("</ul>")
    for name, charts in sections:
        if not charts:
            continue
        parts.append(f"<h2>{name}</h2>")
        parts += [f"<img alt='{png}' src='data:image/png;base64,{b64}'>"
                  for png, b64 in charts]
    if skipped:
        parts.append("<p class='dim'>Not charted: " + ", ".join(skipped) +
                     " — no table with enough numbers in it.</p>")
    return "".join(parts)


def main(job_dir):
    job = Path(job_dir)
    out = job / "out"
    out.mkdir(exist_ok=True)
    inputs = sorted(p for p in (job / "in").iterdir()
                    if p.is_file() and not p.name.startswith("."))
    common.set_state("working")

    sections, findings, skipped = [], [], []
    for path in inputs:
        put_status(job, state="working", step=f"reading {path.name}")
        try:
            tables = tables_from(path)
        except Exception as e:
            skipped.append(f"{path.name} ({type(e).__name__})")
            continue
        if not tables:
            skipped.append(path.name)
            continue
        for name, frame in tables:
            put_status(job, state="working", step=f"charting {name}")
            roles = by_model(name, frame) or by_heuristic(frame)
            charts = []
            try:
                chart_table(name, frame, roles, out, charts, findings)
            except Exception:
                pass          # one bad table must not sink the report
            if charts:
                sections.append((name, charts))
            else:
                skipped.append(name)

    if not sections:
        put_status(job, state="failed",
                   step="nothing chartable arrived — I need a table with numbers "
                        "in it (xlsx, csv, or a Word table)")
        common.set_state("idle")
        return

    title = (json.loads((job / "job.json").read_text(encoding="utf-8")).get("brief") or "").strip() \
        or "Report"
    (out / "report.html").write_text(
        html_report(title, sections, findings, skipped, provenance(job)), encoding="utf-8")
    files = sorted(f.name for f in out.iterdir() if f.is_file() and f.stat().st_size)
    put_status(job, state="done", files=files, empty=[],
               said="\n".join(findings[:10]),
               step=f"done — report.html and {len(files) - 1} chart(s)")
    common.set_state("idle")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("usage: reporter.py <job-dir>")
    try:
        main(sys.argv[1])
    except Exception as e:
        try:
            put_status(Path(sys.argv[1]), state="failed",
                       step=f"the reporter itself failed: {type(e).__name__}: {e}")
        finally:
            common.set_state("idle")
        raise
