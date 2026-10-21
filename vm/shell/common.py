#!/usr/bin/env python3
"""
Shishya — shared bits for the apps that run on the clone's desktop.

Two jobs: talk to the model on the host, and turn any file the user drops in
into plain text.
"""

import json
import os
import re
import subprocess
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path

HOME = Path.home()
# Inside the Lima VM the shared mount lands at ~/shared. On Windows the clone
# is an AppContainer and its home is the container's own storage, so the host
# says where shared/ is rather than the clone guessing — the same folder either
# way, and ~/shared reaches it too, through a junction, for code that builds
# the path itself.
SHARED = Path(os.environ["SHISHYA_SHARED"]) if os.environ.get("SHISHYA_SHARED") \
    else HOME / "shared"
WINDOWS = sys.platform == "win32"


def _sandbox_temp():
    """Temporary folders that work inside the AppContainer.

    Two things broke every tempfile.TemporaryDirectory here — and with it the
    Code app's Run, reading old Office files and the checks. Windows rewrites
    TEMP for a sandboxed process to a folder that does not exist, so Python
    fell back to the current folder; and Python 3.12 makes those folders with
    an ACL naming only the Windows user, which an AppContainer process is not
    — it was locked out of the folder it had just made. So: the real Temp
    (passed as SHISHYA_TEMP), and folders made plainly, inheriting its ACL.
    """
    import secrets
    import tempfile
    real = os.environ.get("SHISHYA_TEMP")
    if not real:
        return
    os.makedirs(real, exist_ok=True)
    os.environ["TEMP"] = os.environ["TMP"] = real
    tempfile.tempdir = real

    def mkdtemp(suffix=None, prefix=None, dir=None):
        base = dir or real
        for _ in range(100):
            path = os.path.join(base, f"{prefix or 'tmp'}{secrets.token_hex(5)}{suffix or ''}")
            try:
                os.mkdir(path)
                return os.path.abspath(path)
            except FileExistsError:
                continue
        raise FileExistsError("no free temporary name")
    tempfile.mkdtemp = mkdtemp


if WINDOWS:
    _sandbox_temp()
WORK = SHARED / "work"
LESSONS = SHARED / "lessons"
TRAINING = SHARED / "training"
TRACES = SHARED / "traces"
STATE = SHARED / ".state"

# The shop's live books, written by the Billing app and read by everything
# else. Plain CSV on purpose: the whole point is that the chat, a job, a
# report and a person with a spreadsheet all read the same two files, and
# a format that needs a library to open is a format somebody cannot audit.
# Who is at the counter. Not a security boundary — the sandbox, the mount
# list and the absence of sudo are the security boundary, and this file sits
# inside the very sandbox it describes. It is what a shop actually needs:
# the person on the till should not be able to void yesterday's takings or
# roll the machine back, and should not have to look at controls that would
# only ever be a mistake. Say "role", never "permission".
ROLE = SHARED / ".role"
# Not ROLES: that name already belongs to the job→model map below.
ROLE_NAMES = ("clerk", "head")


def role():
    """clerk or head. Clerk when unset — the safer of the two to assume."""
    try:
        found = ROLE.read_text(encoding="utf-8").strip().lower()
        return found if found in ROLE_NAMES else "clerk"
    except OSError:
        return "clerk"


def set_role(who):
    who = (who or "").strip().lower()
    if who not in ROLE_NAMES:
        return None
    try:
        SHARED.mkdir(parents=True, exist_ok=True)
        ROLE.write_text(who, encoding="utf-8")
    except OSError:
        pass
    return who


def may(what):
    """Whether the person at the counter may do this.

    Listed rather than inferred, so adding a control means deciding who it is
    for. Everything not named here is allowed to both: the default is the
    ordinary work, and the exceptions are the things that are hard to undo.
    """
    HEAD_ONLY = {"void", "rollback", "egress", "browsing", "snapshot",
                 "schedule", "shell", "models"}
    return role() == "head" or what not in HEAD_ONLY


BILLING = SHARED / "billing"
CATALOGUE = BILLING / "catalogue.csv"     # what is stocked, and at what cost
SALES = BILLING / "sales.csv"             # one row per line item, append-only

# The model runs on the host's GPU and is reached across the sandbox
# boundary. Inside the VM it would fall back to emulated CPU and be far too
# slow to demo.
HOST = os.environ.get("SHISHYA_MODEL_URL", "http://host.lima.internal:11434")

# On Windows the sandbox has no network at all, loopback included, so the
# model is not at an address — it is behind a named pipe whose ACL names this
# container and nobody else (host/broker.py). Same requests, same replies.
PIPE = r"\\.\pipe\shishya-model"
VIA_PIPE = os.environ.get("SHISHYA_MODEL") == "pipe"


class ModelUnreachable(RuntimeError):
    pass


def _pipe_open():
    deadline = time.time() + 15
    while True:
        try:
            return open(PIPE, "r+b", buffering=0)
        except FileNotFoundError:
            raise ModelUnreachable("the model pipe is not there — the Shishya "
                                   "panel is what serves it; is it running?")
        except OSError:
            if time.time() > deadline:          # every instance mid-answer
                raise ModelUnreachable("the model pipe stayed busy for 15 seconds")
            time.sleep(0.1)


def model_request(method, path, body=None, timeout=600):
    """One call to the model service, by whichever route this sandbox has.
    Returns the decoded JSON reply. Raises ModelUnreachable with a sentence a
    person can act on."""
    if not VIA_PIPE:
        data = json.dumps(body).encode() if body is not None else None
        req = urllib.request.Request(f"{HOST}{path}", data, method=method,
                                     headers={"Content-Type": "application/json"})
        try:
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return json.loads(r.read() or b"{}")
        except urllib.error.URLError as e:
            raise ModelUnreachable(
                f"can't reach the model on the host ({e}). If the egress lock is "
                f"on, {HOST} is the one address it should still allow.")

    with _pipe_open() as pipe:
        pipe.write(json.dumps({"method": method, "path": path, "body": body}).encode() + b"\n")
        chunks = []
        while True:
            try:
                chunk = pipe.read(65536)
            except OSError:
                break                               # the broker closed: done
            if not chunk:
                break
            chunks.append(chunk)
    head, _, rest = b"".join(chunks).partition(b"\n")
    status = int(head[:3]) if head[:3].isdigit() else 502
    if status != 200:
        raise ModelUnreachable(head.decode("utf-8", "replace")[4:] or f"the model said {status}")
    return json.loads(rest or b"{}")


def model_stream(method, path, body=None, timeout=3600):
    """Line-by-line replies — a model download reports its progress this way.
    Over the pipe the broker forwards the stream as it arrives."""
    if not VIA_PIPE:
        data = json.dumps(body).encode() if body is not None else None
        req = urllib.request.Request(f"{HOST}{path}", data, method=method,
                                     headers={"Content-Type": "application/json"})
        try:
            with urllib.request.urlopen(req, timeout=timeout) as r:
                for line in r:
                    if line.strip():
                        yield json.loads(line)
        except urllib.error.URLError as e:
            raise ModelUnreachable(str(e))
        return
    with _pipe_open() as pipe:
        pipe.write(json.dumps({"method": method, "path": path, "body": body}).encode() + b"\n")
        buf, head = b"", None
        while True:
            try:
                chunk = pipe.read(65536)
            except OSError:
                break
            if not chunk:
                break
            buf += chunk
            if head is None and b"\n" in buf:
                head, buf = buf.split(b"\n", 1)
                if not head.startswith(b"200"):
                    raise ModelUnreachable(head.decode("utf-8", "replace")[4:])
            while head is not None and b"\n" in buf:
                line, buf = buf.split(b"\n", 1)
                if line.strip():
                    yield json.loads(line)


# The interpreter for running a script as a program of its own. In a windowed
# app on Windows sys.executable is pythonw.exe, which has no console streams;
# python.exe beside it does, and CREATE_NO_WINDOW stops a console flashing up
# for every program the clone runs.
PYTHON = sys.executable
if WINDOWS and PYTHON.lower().endswith("pythonw.exe"):
    PYTHON = PYTHON[:-len("pythonw.exe")] + "python.exe"
NO_WINDOW = 0x08000000 if WINDOWS else 0


def run_python(script, cwd=None, timeout=60):
    """Run a Python file as its own process; returns the CompletedProcess.

    On Windows the child inherits this process's AppContainer token and its
    job, so a generated program is exactly as sandboxed as the app that ran
    it — no network, no route to anything it was not granted.
    """
    return subprocess.run([PYTHON, str(script)], capture_output=True, text=True,
                          encoding="utf-8", errors="replace", timeout=timeout,
                          cwd=str(cwd) if cwd else None, creationflags=NO_WINDOW,
                          env={**os.environ, "PYTHONIOENCODING": "utf-8", "PYTHONUTF8": "1",
                               "MPLBACKEND": "Agg"})
# The fallback when a call names no job. This was deepseek-r1:8b, and the one
# app that forgot to name its jobs — chat, the front door — silently ran a 5GB
# reasoning model while the Models screen showed the 14GB one the owner picked.
# Thirty minutes of "I couldn't write a program for that." The fallback is now
# the owner's model, so forgetting a job name costs nothing but this comment.
MODEL = "gpt-oss:20b"

# ─── one job, one model ─────────────────────────────────────────────────────
# Different work wants different models, and a school should not have to hold
# twenty gigabytes on disk to answer a question about an attendance policy.
# Each job names a model; the Models app writes this file and the apps read it.
ROLES = SHARED / ".models.json"


def model_for(job):
    """Which model does this job. Falls back rather than failing — a missing
    or malformed file must never stop the app from answering."""
    try:
        import json as _json
        chosen = _json.loads(ROLES.read_text(encoding="utf-8")).get(job)
        if chosen:
            return chosen
    except (OSError, ValueError):
        pass
    return MODEL

ACCENT = "#0078d4"


def set_state(state):
    """What the taskbar pill shows: idle | working | wheel | needs."""
    try:
        SHARED.mkdir(parents=True, exist_ok=True)
        STATE.write_text(state, encoding="utf-8")
    except OSError:
        pass


def ask_job(job, prompt, **kw):
    """ask(), told which job it is for, so the right model answers."""
    return ask(prompt, job=job, **kw)


# How long Ollama holds a model in memory after answering. The default is five
# minutes, and the owner's model is 13.8GB — so a question asked six minutes
# after the last one pays twenty to forty seconds of reading that back off
# disk before a single token is generated, with nothing on screen to say why.
# During a demonstration that is the difference between "it thought for a
# moment" and "it has hung".
#
# Fifteen minutes covers a working session and a demonstration without holding
# the memory all day: this Mac has 24GB, the sandbox has 8 of them, and the
# model wants 12.7 — there is not room to be careless about it.
KEEP_ALIVE = "15m"


def warm(job=None):
    """Load the model without asking it anything.

    Ollama treats a generate call with an empty prompt as "load this and
    hold it". Calling it when the sandbox comes up means the first real
    question is answered at speed instead of paying for the load.

    Never raises: the point of warming is that it is optional. If the host's
    Ollama is not up yet, the first real question will load the model itself,
    slowly, which is exactly what happens today.
    """
    try:
        model_request("POST", "/api/generate",
                      {"model": model_for(job) if job else MODEL,
                       "prompt": "", "keep_alive": KEEP_ALIVE}, timeout=300)
        return True
    except Exception:
        return False


def ask(prompt, system=None, schema=None, timeout=600, budget=None, job=None, model=None):
    body = {
        # model= names one outright (the benchmark); otherwise the job's.
        "model": model or (model_for(job) if job else MODEL),
        "prompt": prompt,
        "stream": False,
        "think": False,
        "keep_alive": KEEP_ALIVE,
        "options": {"temperature": 0.2, "num_ctx": 8192, "num_predict": 2048},
    }
    if system:
        body["system"] = system
    if schema:
        body["format"] = schema
    def once(budget):
        body["options"]["num_predict"] = budget
        # num_ctx has to hold the prompt AND the answer, and it was a fixed
        # 8,192 while the code path asked for 6,144 tokens of answer. A 2,700
        # token file listing then left the prompt truncated, so the model was
        # told to write a program from instructions that had been cut off —
        # and answered "I couldn't write a program for that", correctly.
        #
        # Worse, the empty-answer retry below doubles the budget, which made
        # the room for the prompt negative and guaranteed the second attempt
        # failed too. Both attempts ran to completion first: half an hour of
        # a 20B model reasoning about a question it could not see.
        #
        # Sized per call instead. Three characters per token is deliberately
        # pessimistic — column names and code tokenise worse than prose.
        need = (len(prompt) + len(system or "")) // 3 + budget + 512
        body["options"]["num_ctx"] = max(8192, min(need, 32768))
        return _strip_think(model_request("POST", "/api/generate", body,
                                          timeout=timeout).get("response", ""))

    def attempt(budget):
        try:
            return once(budget)
        except ModelUnreachable:
            # Ollama's runner can die while it swaps one large model off a
            # 16 GB graphics card for another (gpt-oss's 13 GB out, the
            # answers model in): "llama-server process has terminated". The
            # next request starts a fresh runner, so one more try after a
            # pause goes through instead of failing the person's question.
            time.sleep(5)
            return once(budget)

    if budget:
        # Code needs more room than prose: the model reasons first, and a
        # script that stops halfway is worse than no script.
        body["options"]["num_predict"] = budget
    try:
        answer = attempt(body["options"]["num_predict"])
        if not answer:
            # It thought until it ran out of room. Once more, with more room.
            answer = attempt(body["options"]["num_predict"] * 2)
        return answer or (
            "I thought my way through that without reaching an answer — it was "
            "too broad to settle. Ask about one thing at a time and I'll get "
            "there. If you want a number counted across the whole sheet, ask "
            "me to work it out and I'll write a program instead.")
    except ModelUnreachable as e:
        raise RuntimeError(str(e))



def _strip_think(text):
    """Remove the model's reasoning, or report that there was no answer.

    deepseek-r1 emits <think>...</think> even with thinking disabled. If it
    spends its whole token budget reasoning the block is never closed, and
    there is no answer in the reply at all — the earlier version fell back to
    returning the raw text, so the user saw the reasoning presented as the
    answer. Returning empty lets the caller retry with more room instead.
    """
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.S).strip()
    if "<think>" in text:
        return ""                     # opened and never closed: truncated
    # qwen3's template opens the block itself, in the prompt, so the reply
    # carries only the close: the reasoning, then </think>, then the answer.
    # Without this the reasoning reached the user as the answer.
    if "</think>" in text:
        text = text.rsplit("</think>", 1)[1].strip()
    return text


PLAIN = {".txt", ".md", ".csv", ".tsv", ".json", ".log", ".py", ".xml", ".html"}

# Read Office files with the libraries that understand them rather than by
# shelling out to LibreOffice. `--headless --cat` silently produces nothing
# for spreadsheets, so an .xlsx a person can plainly open came back as
# "LibreOffice could not read this". These libraries are already installed
# for writing documents; using them to read is faster, needs no subprocess,
# and cannot be wedged by a LibreOffice instance that is already running.
NATIVE = {".xlsx": "sheet", ".xlsm": "sheet", ".docx": "doc", ".pptx": "deck"}

# Older and OpenDocument formats still need LibreOffice, converting to a
# format we can read rather than asking it to print to stdout.
CONVERT = {".odt", ".ods", ".odp", ".doc", ".xls", ".ppt", ".rtf"}


def extract(path):
    """Returns (text, how). `how` is shown in the UI so the route a file took
    is visible — and so an empty result explains itself rather than looking
    like the file was ignored."""
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix in PLAIN:
        return path.read_text(encoding="utf-8", errors="replace"), "read as text"

    if suffix == ".pdf":
        out = _run(["pdftotext", "-layout", str(path), "-"])
        if out.strip():
            return out, "extracted with pdftotext"
        # No poppler on Windows. pypdf reads the text layer in-process, which
        # is all pdftotext was doing for a PDF that has one.
        try:
            from pypdf import PdfReader
            out = "\n\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)
            if out.strip():
                return out, "read with pypdf"
        except Exception:
            pass
        if WINDOWS:
            return "", ("this PDF has no text layer — it is a scan, and reading "
                        "scans needs OCR, which the Windows build does not include yet")
        # A scan. OCR it — tesseract went in on 23 Aug 26, with Hindi and
        # Telugu alongside English, because that is what actually lands on a
        # desk in Hyderabad. Slow (about a second a page), so only reached
        # when the text layer is empty.
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            side = Path(tmp) / "ocr.txt"
            _run(["ocrmypdf", "--sidecar", str(side), "-l", "eng+hin+tel",
                  "--force-ocr", "--output-type", "none", str(path), "-"],
                 timeout=300)
            try:
                text = side.read_text(encoding="utf-8", errors="replace")
            except OSError:
                text = ""
        if text.strip():
            return text, "scanned PDF — read with OCR"
        return "", "PDF has no text layer and OCR found nothing legible"

    if suffix in NATIVE:
        try:
            reader = {"sheet": _read_sheet, "doc": _read_doc, "deck": _read_deck}
            text = reader[NATIVE[suffix]](path)
            if text.strip():
                return text, f"read {suffix[1:]} directly"
            return "", f"{suffix[1:]} opened but is empty"
        except Exception as e:
            return "", f"could not read {suffix[1:]}: {type(e).__name__}"

    if suffix in CONVERT:
        return _via_libreoffice(path)

    try:
        raw = path.read_bytes()[:200_000]
        if b"\x00" not in raw[:2000]:
            return raw.decode("utf-8", errors="replace"), f"unknown type {suffix}, read as text"
    except OSError:
        pass
    return "", f"can't read {suffix or 'this file'}"


def _read_sheet(path):
    """Every sheet as pipe-separated rows, with the sheet named.

    Pipes rather than commas because real cells contain commas, and a model
    reading `12,500 | 3 | yes` should not have to guess where the columns
    are. Blank trailing cells are dropped so short rows stay short.
    """
    from openpyxl import load_workbook
    # data_only: we want what the sheet shows, not "=SUM(B2:B9)".
    wb = load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            while cells and cells[-1] == "":
                cells.pop()
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            out.append(f"## Sheet: {ws.title}\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(out)


def _read_doc(path):
    from docx import Document
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_deck(path):
    from pptx import Presentation
    deck = Presentation(str(path))
    parts = []
    for i, slide in enumerate(deck.slides, 1):
        lines = [f"## Slide {i}"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    text = "".join(r.text for r in p.runs).strip()
                    if text:
                        lines.append(text)
        if len(lines) > 1:
            parts.append("\n".join(lines))
    return "\n\n".join(parts)


def _via_libreoffice(path):
    """Convert to something readable in a temp directory. Never --cat: it
    prints nothing for spreadsheets and fails silently."""
    import tempfile
    target = "csv" if path.suffix.lower() in {".ods", ".xls"} else "txt:Text"
    office = libreoffice()
    if not office:
        return "", f"{path.suffix} needs LibreOffice to read, and it is not installed"
    with tempfile.TemporaryDirectory() as tmp:
        _run([office, "--headless", "--convert-to", target,
              "--outdir", tmp, str(path)], timeout=150)
        for made in Path(tmp).iterdir():
            body = made.read_text(encoding="utf-8", errors="replace")
            if body.strip():
                return body, f"converted with LibreOffice to {made.suffix[1:]}"
    return "", "LibreOffice could not convert this"


def libreoffice():
    """The LibreOffice program, or None. On Windows it is not on PATH — it
    sits in Program Files, which every AppContainer may read and run."""
    import shutil
    found = shutil.which("libreoffice") or shutil.which("soffice")
    if found:
        return found
    if WINDOWS:
        for root in (os.environ.get("ProgramFiles", r"C:\Program Files"),
                     os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")):
            candidate = Path(root) / "LibreOffice" / "program" / "soffice.exe"
            if candidate.exists():
                return str(candidate)
    return None


def _run(cmd, timeout=60):
    try:
        p = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout,
                           encoding="utf-8", errors="replace", creationflags=NO_WINDOW)
        return p.stdout
    except (subprocess.SubprocessError, OSError):
        return ""


def chunk(text, size=1400, overlap=180):
    """Overlapping windows so a fact sitting on a boundary is not cut in half."""
    text = " ".join(text.split())
    if len(text) <= size:
        return [text] if text else []
    out, i = [], 0
    while i < len(text):
        out.append(text[i:i + size])
        i += size - overlap
    return out
