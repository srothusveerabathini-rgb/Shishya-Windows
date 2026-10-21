#!/usr/bin/env python3
"""
Shishya — the machinery for answering with a program instead of by eye.

Describing the material so a model can write code against it, repairing the
slips that model reliably makes, and running what it wrote. Split out of
chat_app.py, where it grew, because the worker, Findings, the benchmark and
the Windows build's Qt apps all trust the same machinery — and none of them
should have to import a GTK window to reach it.
"""

import re
import subprocess
import tempfile
from pathlib import Path

import common


RUN_TIMEOUT = 60


DECIDE = {
    "type": "object",
    "properties": {
        "action": {"type": "string",
                   "enum": ["answer", "spreadsheet", "document", "deck", "code", "clarify"]},
        "because": {"type": "string"},
        "brief": {"type": "string"},
    },
    "required": ["action", "because", "brief"],
}


DECIDE_SYSTEM = """You decide what to do with a request about the user's own documents.

Choose exactly one:
  answer      — any question about rules, policy, eligibility, or what a
                particular record says. This is the default for questions.
                Use it even when the question describes a specific case the
                documents do not name — working out what the rules say about
                a case IS answering.
  spreadsheet — they asked for a table or data file (.xlsx)
  document    — they asked for prose: a note, letter, report, brief (.docx)
  deck        — they asked for slides (.pptx)
  code        — the answer is a number that has to be counted, totalled,
                filtered or ranked across many rows. You are unreliable doing
                that by eye, so write a program instead. Looking up one record
                is not this — that is answer.
  clarify     — you genuinely cannot tell what they are asking for. Almost
                never the right choice. If you could make a reasonable
                attempt, make it. Never use clarify merely because the
                documents might not contain the answer; try, and say so if
                they do not.

"because" is one short sentence to the user saying what you are about to do.
"brief" restates the job precisely for whoever carries it out.
Answer as JSON only."""


# Which operating system the script will run on is said, not assumed: a
# model told "Linux" on Windows writes '/tmp/out.xlsx', which does not exist
# there and would not be inside the sandbox if it did.
PLATFORM = "Windows" if common.WINDOWS else "Linux"

CODE_SYSTEM = (
    "You write short, complete Python 3 scripts. Output code only — no "
    f"explanation, no markdown fences. Python 3 on {PLATFORM} with pandas, "
    "openpyxl, python-docx and python-pptx available. Build paths with "
    "pathlib — Path.home() / 'shared' / 'training' / '<name>' — never a "
    "string starting with ~. Print the result. Under 40 lines."
)


def _header_row(ws, look=14):
    """Which row holds the column names. Not always the first.

    A sheet built for a person to read opens with a title, a note about what
    it is, and a block of assumptions; the columns start ten rows down. Taking
    row 1 regardless named this workbook's margin columns "PRODUCT-WISE GROSS
    MARGIN & MOVEMENT ANALYSIS" and eighteen blanks, so every question about
    margin was answered against a sheet whose real columns were never seen.

    The header is the widest row near the top — a title row has one cell, a
    header has nineteen. Ties go to the earliest, which keeps a genuine
    one-column sheet working.
    """
    best, widest = 1, 0
    for n, row in enumerate(ws.iter_rows(min_row=1, max_row=look, values_only=True), 1):
        filled = sum(1 for c in row if c is not None and str(c).strip())
        if filled > widest:
            best, widest = n, filled
    return best


def _describe_sheet(ws, max_rows, max_cols):
    """One sheet's columns, and the values a short column actually contains.

    Names alone are not enough. Told the column was called "Stock Status",
    the model filtered for "Low" when every row says "Low Stock", and
    reported that nothing was low on stock across a sheet with 149 such
    rows — wrong, confident, and impossible to spot without opening the file.
    Columns with few distinct values get those values listed; a column of
    1,000 different prices obviously does not.
    """
    at = _header_row(ws)
    rows = ws.iter_rows(min_row=at, values_only=True)
    try:
        raw_header = [c for c in next(rows) if c is not None]
    except StopIteration:
        return None
    if not raw_header:
        return None
    header = [str(c) for c in raw_header][:max_cols]

    # Which headers are real dates rather than text. Printed as strings they
    # look like text — "2025-04-01 00:00:00" — and the model kept writing
    # .startswith('2025') against them, which fails with "datetime.datetime
    # object has no attribute startswith". Four of eight test questions died
    # on exactly that. Saying so plainly fixes all of them.
    import datetime as _dt
    dated = [str(c) for c in raw_header if isinstance(c, (_dt.datetime, _dt.date))]

    seen = [set() for _ in header]
    counted = 0
    for row in rows:
        if counted >= max_rows:
            break
        counted += 1
        for i, cell in enumerate(row[:len(header)]):
            if cell is not None and len(seen[i]) <= 13:
                seen[i].add(str(cell))

    out = [f"      columns: {', '.join(header)}"]
    if len(raw_header) > max_cols:
        out.append(f"      ({len(raw_header) - max_cols} further columns not listed)")
    if at > 1:
        out.append(f"      NOTE: the column names are on row {at}, not row 1 — "
                   f"read it with pd.read_excel(..., sheet_name={ws.title!r}, "
                   f"header={at - 1}). Rows above that are a title and notes.")
    if dated:
        out.append(
            f"      NOTE: {len(dated)} of those column headers are real datetime "
            f"objects, not text — {dated[0]} through {dated[-1]}. Select them "
            f"with [c for c in df.columns if isinstance(c, datetime)] after "
            f"`from datetime import datetime`. String methods such as "
            f"startswith or str.contains will raise on them. When filtering "
            f"columns BY NAME, guard first — "
            f"[c for c in df.columns if isinstance(c, str) and c.startswith(...)] "
            f"— or the datetime headers will raise as soon as you reach them.")

    # Describe repeated column families rather than listing them one by one.
    # This sheet keeps suppliers wide — "Purchased From 1".."Purchased From 10"
    # each paired with a "Purchase Qty N" and a "Net Purchase Price N". Left to
    # infer that from a flat list of 54 names, the model wrote nothing at all
    # when asked which suppliers we buy most from. Said plainly, it has a
    # chance: the answer needs those pairs stacked before they can be grouped.
    families = {}
    for name in header:
        stem = re.sub(r"\s*\d+$", "", name)
        if stem != name:
            families.setdefault(stem, []).append(name)
    for stem, cols in families.items():
        if len(cols) > 2:
            out.append(f"      {stem} 1..{len(cols)} — one column per repeat")
    if len(families) > 1 and any(len(c) > 2 for c in families.values()):
        pairs = " and ".join(f'"{stem} N"' for stem in families if len(families[stem]) > 2)
        out.append(f"      these repeat together: {pairs} share the same N, so "
                   f"to total by value across them, stack the pairs first "
                   f"(pd.concat of each N, or pd.wide_to_long) and then group")

    for name, values in zip(header, seen):
        if 1 < len(values) <= 12:
            out.append(f"      {name} contains only: {', '.join(sorted(values))}")
    return "\n".join(out)


def sheet_shape(path, max_rows=400, max_sheets=16, max_cols=44):
    """Every sheet in the workbook, not only the first.

    A one-sheet export was the only shape this had ever been given. A real
    workbook is a stack of sheets that refer to each other — a master list,
    and a dozen views built on top of it — and describing only the first left
    the model computing from raw columns while a sheet answering the exact
    question sat three tabs away, unmentioned.
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    blocks = []
    for ws in wb.worksheets[:max_sheets]:
        try:
            block = _describe_sheet(ws, max_rows, max_cols)
        except Exception:
            continue          # one unreadable sheet must not hide the rest
        if block:
            blocks.append(f"    sheet {ws.title!r}:\n{block}")
    extra = len(wb.worksheets) - max_sheets
    wb.close()
    if not blocks:
        return None
    if extra > 0:
        blocks.append(f"    ({extra} further sheets not described)")
    return "\n".join(blocks)


# The listing goes into every code prompt, so it cannot grow without limit —
# one fourteen-sheet workbook took it past the context window on its own. The
# window is now sized to fit (see common.ask), but a prompt that grows with the
# training folder is one that will eventually be mostly file listing.
#
# Trimmed, never dropped. The first version skipped any file that would not fit
# whole, which threw away the largest workbook — the only one that could answer
# the question — while keeping four small files nobody had asked about. A path
# always goes in, because the model cannot read a file it was not told about;
# only the sheet detail is cut, at a line boundary, and it says so.
LISTING_BUDGET = 16_000


NOTE_BUDGET = 3_500


def describe_billing():
    """The shop's own books, described so a program can read them.

    Kept out of the material folder deliberately: training material is
    reference — a policy, a price list, a manual — and these two files are
    live operational data that changed while you were reading this. They are
    named separately in the prompt for the same reason, so the model is never
    guessing whether a row is history or today.
    """
    import csv as _csv
    out = []
    for path, what in ((common.CATALOGUE, "what is stocked right now, with cost, "
                                          "price, stock on hand and reorder point"),
                       (common.SALES, "every line item ever sold, append-only, one "
                                      "row per product per bill")):
        if not path.exists():
            continue
        try:
            with open(path, newline="", encoding="utf-8") as f:
                reader = _csv.reader(f)
                header = next(reader, [])
                rows = sum(1 for _ in reader)
        except OSError:
            continue
        out.append(f"  {path}\n"
                   f"      {what}\n"
                   f"      columns: {', '.join(header)}   ({rows:,} rows)")
    if not out:
        return ""
    return ("\n\nTHE SHOP'S LIVE BOOKS — written by the Billing app, current as "
            "of now. `at` is 'YYYY-MM-DD HH:MM:SS' text; parse with "
            "pd.to_datetime. `profit` is already (price − cost) × qty at the "
            "cost on the day of sale — sum it, never recompute margin from "
            "today's catalogue cost:\n" + "\n".join(out))


def csv_shape(path, look=500, few=12):
    """A CSV's columns, row count, and what its short columns contain.

    Only spreadsheets used to be described. Asked to count the students in
    class-9b-term1.csv, every model saw the file's name and no columns, and
    borrowed "Students" from the workbook listed beside it.
    """
    import csv
    with open(path, encoding="utf-8-sig", errors="replace", newline="") as fh:
        rows = list(csv.reader(fh, delimiter="\t" if path.suffix.lower() == ".tsv" else ","))
    if not rows:
        return None
    header, body = rows[0], rows[1:]
    out = [f"    columns: {', '.join(header)}", f"    {len(body)} rows"]
    for i, col in enumerate(header):
        seen = {r[i].strip() for r in body[:look] if i < len(r) and r[i].strip()}
        if 1 < len(seen) <= few and not all(v.replace(".", "", 1).isdigit() for v in seen):
            out.append(f"    {col} contains only: {', '.join(sorted(seen))}")
    return "\n".join(out)


def describe_material():
    lines, spent = [], 0
    for p in sorted(common.TRAINING.iterdir()):
        if not p.is_file() or p.name.startswith("."):
            continue
        entry = f"  {common.TRAINING / p.name}"
        spent += len(entry)
        shape = None
        if p.suffix.lower() in (".xlsx", ".xlsm"):
            try:
                shape = sheet_shape(p)
            except Exception:
                shape = None
        elif p.suffix.lower() in (".csv", ".tsv"):
            try:
                shape = csv_shape(p)
            except Exception:
                shape = None
        if shape:
            room = LISTING_BUDGET - spent
            if len(shape) > room:
                kept = shape[:max(room, 0)].rsplit("\n", 1)[0]
                shape = (kept + "\n    (further sheets not described)") if kept else None
            if shape:
                entry += "\n" + shape
                spent += len(shape)
        lines.append(entry)
    if not lines:
        return "There are no files yet."
    listing = ("These files exist — use these exact paths and column names:\n"
               + "\n".join(lines))

    # Written guidance reaches the code path too, not only the answering one.
    #
    # A note in the material saying "this sheet costs every sale at the
    # cheapest of ten purchase prices, so its margins are three times too
    # good" is exactly what someone writing a program needs, and it was only
    # ever retrieved for prose questions. The program went on computing from
    # a column the note exists to warn about. A file named *-notes.md is the
    # convention: material about the material.
    # Every notes file that fits, not the first two. A count cap meant the
    # third file written was silently never read — and nothing says so.
    notes, spent = [], 0
    for p in sorted(common.TRAINING.glob("*notes.md")):
        try:
            text = p.read_text(encoding="utf-8")[:NOTE_BUDGET]
        except OSError:
            continue
        if spent + len(text) > 3 * NOTE_BUDGET:
            break
        notes.append(f"--- {p.name} ---\n{text}")
        spent += len(text)
    if notes:
        listing += ("\n\nRead these notes before writing the program. They "
                    "describe traps in the data above and override anything a "
                    "column name implies:\n" + "\n".join(notes))
    listing += describe_billing()
    return listing


def repair_regex(code):
    r"""Make string literals containing regex escapes raw.

    Asked which supplier gives the best average price, it reached for
    pd.wide_to_long with suffix='\d+' — no r prefix — which Python reads as
    an invalid escape and pandas then fails on. The intent is unambiguous
    every time: a backslash followed by a regex class letter inside a short
    literal is a pattern, not text.
    """
    import re

    def raw(m):
        quote, body = m.group(1), m.group(2)
        return f"r{quote}{body}{quote}"

    # Short literals only, and only those carrying a regex escape. A long
    # string with a stray backslash is more likely to be prose.
    return re.sub(r"(?<![rbRB])(['\"])((?:[^'\"\\]|\\.){0,40}?\\[dDwWsSbBAZ][^'\"]{0,20}?)\1",
                  raw, code)


def repair(raw):
    """Everything we fix about generated code, in one place."""
    return repair_imports(repair_regex(repair_calls(repair_paths(strip_fences(raw)))))


# Names the models use without importing them. The commonest single failure
# on the code benchmark — 72 of 95 wrong scripts — was Path(...) with no
# `from pathlib import Path`, a script right in every other way.
IMPORTS = [
    (re.compile(r"\bPath\s*[.(]"), re.compile(r"^\s*from\s+pathlib\s+import\s+.*\bPath\b", re.M),
     "from pathlib import Path"),
    (re.compile(r"\bpd\."), re.compile(r"^\s*import\s+pandas\s+as\s+pd\b", re.M), "import pandas as pd"),
    (re.compile(r"\bnp\."), re.compile(r"^\s*import\s+numpy\s+as\s+np\b", re.M), "import numpy as np"),
    (re.compile(r"\bplt\."), re.compile(r"^\s*(import\s+matplotlib\.pyplot\s+as\s+plt|from\s+matplotlib\s+import\s+pyplot\s+as\s+plt)\b", re.M),
     "import matplotlib\nmatplotlib.use('Agg')\nimport matplotlib.pyplot as plt"),
]
STDLIB = ("re", "json", "os", "math", "csv", "datetime", "statistics", "collections")


def repair_imports(code):
    """Add an import the script uses but never made."""
    missing = [line for used, made, line in IMPORTS if used.search(code) and not made.search(code)]
    for mod in STDLIB:
        if re.search(rf"\b{mod}\.", code) and not re.search(
                rf"^\s*(import\s+[\w, ]*\b{mod}\b|from\s+{mod}\s+import)", code, re.M) \
                and not re.search(rf"^\s*from\s+\w+\s+import\s+.*\b{mod}\b", code, re.M):
            missing.append(f"import {mod}")
    return "\n".join(missing + [code]) if missing else code


def usable(code):
    """Is this a script, or the model's apology?"""
    return bool(code.strip()) and not code.lstrip().startswith(
        ("I ran out of room", "I thought my way", "I couldn't"))


def run_script(code):
    """Run it in the sandbox and return what it said. Errors come back whole —
    the traceback is the useful part when the model has to read it back."""
    with tempfile.TemporaryDirectory() as tmp:
        script = Path(tmp) / "job.py"
        script.write_text(code, encoding="utf-8")
        try:
            p = common.run_python(script, cwd=Path.home(), timeout=RUN_TIMEOUT)
            return (p.stdout or "").strip(), (p.stderr or "").strip()
        except subprocess.TimeoutExpired:
            return "", f"it was still running after {RUN_TIMEOUT}s"


# Methods this model writes without their brackets. Shown the traceback it
# made the identical slip a second time, so this is repaired mechanically
# rather than explained again: `df.loc[df['x'].idxmin]` hands pandas the
# method instead of the number and dies with "unhashable type: DataFrame".
#
# Only methods. Properties — size, shape, values, empty, columns, index,
# dtypes, T — are correct without brackets and adding them would break code
# that was already right.
NEEDS_BRACKETS = (
    "idxmin", "idxmax", "mean", "median", "nunique", "unique", "tolist",
    "to_dict", "dropna", "drop_duplicates", "reset_index", "value_counts",
    "sort_values", "sort_index", "count", "sum", "min", "max", "abs",
    "round", "first_valid_index", "last_valid_index", "keys", "item",
)


def repair_calls(code):
    """Add the brackets the model left off a method call."""
    import re
    pattern = r"\.(" + "|".join(NEEDS_BRACKETS) + r")\b(?!\s*[\(\w=])"
    return re.sub(pattern, r".\1()", code)


def repair_paths(code):
    """Point filenames at files that exist.

    Given "01_20Retail_20Inventory.xlsx" the model wrote
    "01_20Retail_Inventory.xlsx" — one fragment dropped. The script then fails
    on a missing file, which reads as the whole tool being broken rather than
    as a typo. Any quoted name ending in a document extension that does not
    exist is matched against the real ones and replaced when there is a clear
    best match.
    """
    import difflib
    import re

    real = [p.name for p in common.TRAINING.iterdir() if p.is_file()]
    real += [p.name for p in common.WORK.iterdir() if p.is_file()] if common.WORK.exists() else []
    if not real:
        return code

    def swap(m):
        quote, name = m.group(1), m.group(2)
        if name in real:
            return m.group(0)
        close = difflib.get_close_matches(name, real, n=1, cutoff=0.6)
        return f"{quote}{close[0]}{quote}" if close else m.group(0)

    return re.sub(r"(['\"])([^'\"/]+\.(?:xlsx|xlsm|csv|docx|pptx|pdf|txt|md))\1",
                  swap, code)


def strip_fences(text):
    import re
    m = re.search(r"```(?:python)?\s*([\s\S]*?)```", text)
    return (m.group(1) if m else text).strip()
